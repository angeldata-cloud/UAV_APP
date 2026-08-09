#!/usr/bin/env python3
"""驗證 4 份正式 Excel，並安全產生 App 題庫與 stable ID manifest。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "source-excel"
OUTPUT_PATH = ROOT / "public" / "data" / "questions.json"
MAP_PATH = ROOT / "data" / "question-id-map.json"
REPORT_PATH = ROOT / "data" / "question-bank-report.json"
SHEET_NAME = "dB題庫"


@dataclass(frozen=True)
class Source:
    chapter: str
    title: str
    filename: str


SOURCES = (
    Source("chapter-01", "第一章：民用航空法及相關法規", "01-讀給你聽-UAV普科題庫_第一章：民用航空法及相關法規.xlsx"),
    Source("chapter-02", "第二章：基礎飛行原理", "02-讀給你聽-UAV普科題庫_第二章：基礎飛行原理.xlsx"),
    Source("chapter-03", "第三章：氣象", "03-讀給你聽-UAV普科題庫_第三章：氣象.xlsx"),
    Source("chapter-04", "第四章：緊急處置與飛行決策", "04-讀給你聽-UAV普科題庫_第四章：緊急處置與飛行決策.xlsx"),
)
REQUIRED_HEADERS = ("題號", "題目", "A", "B", "C", "D", "答案", "題解")


class ValidationErrors(Exception):
    def __init__(self, errors: list[str]):
        super().__init__("\n".join(errors))
        self.errors = errors


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized_header(value: Any) -> str:
    return text(value).rstrip(".．。")


def normalized_answer(value: Any) -> str | None:
    raw = text(value).upper()
    match = re.fullmatch(r"([ABCD])[.．。]?", raw)
    return match.group(1) if match else None


def canonical_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def digest(value: Any) -> str:
    return hashlib.sha256(canonical_json(value).encode("utf-8")).hexdigest()


def read_excel_sources(source_dir: Path = SOURCE_DIR) -> list[dict[str, Any]]:
    errors: list[str] = []
    expected_names = {source.filename for source in SOURCES}
    actual_names = {
        path.name
        for path in source_dir.glob("*.xlsx")
        if not path.name.startswith("~$")
    } if source_dir.exists() else set()

    for name in sorted(expected_names - actual_names):
        errors.append(f"缺少正式 Excel：{name}")
    for name in sorted(actual_names - expected_names):
        errors.append(f"source-excel 中有未預期的 Excel：{name}。請只保留正式的 4 份 Excel。")

    questions: list[dict[str, Any]] = []
    seen_source_keys: dict[str, tuple[str, int]] = {}
    if errors:
        raise ValidationErrors(errors)

    for source in SOURCES:
        path = source_dir / source.filename
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # pragma: no cover - openpyxl supplies varying detail
            errors.append(f"{source.filename}：無法開啟 Excel（{exc}）。")
            continue

        if SHEET_NAME not in workbook.sheetnames:
            errors.append(
                f"{source.filename}：缺少必要工作表「{SHEET_NAME}」；目前工作表為："
                f"{ '、'.join(workbook.sheetnames) or '無' }。"
            )
            workbook.close()
            continue

        sheet = workbook[SHEET_NAME]
        raw_headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [normalized_header(value) for value in raw_headers]
        positions = {header: index for index, header in enumerate(headers) if header}
        missing = [header for header in REQUIRED_HEADERS if header not in positions]
        if missing:
            errors.append(
                f"{source.filename}，工作表「{SHEET_NAME}」：缺少必要欄位「{'、'.join(missing)}」。"
            )
            workbook.close()
            continue

        for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            row = {header: text(values[index] if index < len(values) else None) for header, index in positions.items()}
            if not any(row.get(header, "") for header in REQUIRED_HEADERS):
                continue

            for header in ("題號", "題目", "A", "B", "C", "D"):
                if not row.get(header, ""):
                    errors.append(f"{source.filename}，第 {excel_row} 列：欄位「{header}」不可為空。")

            answer = normalized_answer(row.get("答案"))
            if answer is None:
                shown = row.get("答案", "") or "空白"
                errors.append(
                    f"{source.filename}，第 {excel_row} 列：答案為「{shown}」，每題只能有一個答案 A、B、C 或 D。"
                )

            source_no = row.get("題號", "")
            source_key = f"{source.chapter}:{source_no}"
            if source_key in seen_source_keys:
                previous_file, previous_row = seen_source_keys[source_key]
                errors.append(
                    f"{source.filename}，第 {excel_row} 列：題號「{source_no}」重複；"
                    f"第一次出現在 {previous_file} 第 {previous_row} 列。"
                )
            else:
                seen_source_keys[source_key] = (source.filename, excel_row)

            questions.append(
                {
                    "sourceKey": source_key,
                    "sourceChapter": source.chapter,
                    "sourceChapterTitle": source.title,
                    "sourceQuestionNo": source_no,
                    "question": row.get("題目", ""),
                    "options": {letter: row.get(letter, "") for letter in "ABCD"},
                    "answer": answer or row.get("答案", ""),
                    "explanation": row.get("題解", ""),
                    "_sourceFile": source.filename,
                    "_sourceRow": excel_row,
                }
            )
        workbook.close()

    if errors:
        raise ValidationErrors(errors)
    return questions


def load_mapping(path: Path = MAP_PATH) -> dict[str, Any]:
    if not path.exists():
        raise ValidationErrors([f"缺少 stable ID 對照檔：{path.relative_to(ROOT)}。為保護既有學習紀錄，已停止產生題庫。"])
    try:
        mapping = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValidationErrors([f"stable ID 對照檔無法讀取：{exc}"]) from exc

    entries = mapping.get("questions")
    next_id = mapping.get("nextStableQuestionId")
    if mapping.get("schemaVersion") != 1 or not isinstance(entries, dict) or not isinstance(next_id, int):
        raise ValidationErrors(["stable ID 對照檔格式錯誤，必須包含 schemaVersion、nextStableQuestionId 與 questions。"])

    ids = [entry.get("stableQuestionId") for entry in entries.values() if isinstance(entry, dict)]
    if len(ids) != len(entries) or any(not isinstance(value, int) or value < 1 for value in ids):
        raise ValidationErrors(["stable ID 對照檔含有無效的 stableQuestionId。"])
    if len(ids) != len(set(ids)):
        raise ValidationErrors(["stable ID 對照檔含有重複的 stableQuestionId。"])
    if ids and next_id <= max(ids):
        raise ValidationErrors(["stable ID 對照檔的 nextStableQuestionId 必須大於所有既有 ID。"])
    return mapping


def load_previous_questions(path: Path = OUTPUT_PATH) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValidationErrors([f"現有 questions.json 無法讀取：{exc}"]) from exc
    questions = data.get("questions")
    if not isinstance(questions, list):
        raise ValidationErrors(["現有 questions.json 缺少 questions 陣列。"])
    return questions


def assign_stable_ids(
    raw_questions: list[dict[str, Any]], mapping: dict[str, Any]
) -> tuple[list[dict[str, Any]], list[str]]:
    entries = mapping["questions"]
    next_id = mapping["nextStableQuestionId"]
    new_source_keys: list[str] = []
    assigned: list[dict[str, Any]] = []

    for raw in raw_questions:
        source_key = raw["sourceKey"]
        entry = entries.get(source_key)
        if entry is None:
            entry = {"stableQuestionId": next_id}
            entries[source_key] = entry
            next_id += 1
            new_source_keys.append(source_key)
        stable_id = entry["stableQuestionId"]
        question = {key: value for key, value in raw.items() if not key.startswith("_")}
        question["id"] = stable_id
        question["stableQuestionId"] = stable_id
        revision_input = {key: value for key, value in question.items() if key not in ("id", "stableQuestionId", "contentRevision")}
        question["contentRevision"] = digest(revision_input)
        assigned.append(question)

    mapping["nextStableQuestionId"] = next_id
    if len({question["stableQuestionId"] for question in assigned}) != len(assigned):
        raise ValidationErrors(["產生結果含有重複的 stableQuestionId，已停止更新。"])
    return assigned, new_source_keys


def comparable(question: dict[str, Any]) -> dict[str, Any]:
    return {
        "question": question.get("question", ""),
        "options": question.get("options", {}),
        "answer": normalized_answer(question.get("answer")) or question.get("answer", ""),
        "explanation": question.get("explanation", ""),
    }


def build_report(
    previous: list[dict[str, Any]],
    questions: list[dict[str, Any]],
    new_source_keys: list[str],
    mapping: dict[str, Any],
) -> dict[str, Any]:
    previous_by_id = {question.get("stableQuestionId", question.get("id")): question for question in previous}
    current_by_id = {question["stableQuestionId"]: question for question in questions}
    updated: list[dict[str, Any]] = []
    answer_changed: list[dict[str, Any]] = []

    for stable_id in sorted(previous_by_id.keys() & current_by_id.keys()):
        before = previous_by_id[stable_id]
        after = current_by_id[stable_id]
        if comparable(before) == comparable(after):
            continue
        fields: list[str] = []
        if before.get("question", "") != after["question"]:
            fields.append("題目")
        for letter in "ABCD":
            if before.get("options", {}).get(letter, "") != after["options"][letter]:
                fields.append(f"選項{letter}")
        old_answer = normalized_answer(before.get("answer")) or before.get("answer", "")
        if old_answer != after["answer"]:
            fields.append("答案")
            answer_changed.append(
                {
                    "stableQuestionId": stable_id,
                    "sourceKey": after["sourceKey"],
                    "oldAnswer": old_answer,
                    "newAnswer": after["answer"],
                }
            )
        if before.get("explanation", "") != after["explanation"]:
            fields.append("題解")
        updated.append({"stableQuestionId": stable_id, "sourceKey": after["sourceKey"], "fields": fields})

    removed_ids = sorted(previous_by_id.keys() - current_by_id.keys())
    removed = []
    source_key_by_id = {
        entry["stableQuestionId"]: source_key for source_key, entry in mapping["questions"].items()
    }
    for stable_id in removed_ids:
        removed.append({"stableQuestionId": stable_id, "sourceKey": source_key_by_id.get(stable_id, "unknown")})

    added = [
        {"stableQuestionId": question["stableQuestionId"], "sourceKey": question["sourceKey"]}
        for question in questions
        if question["stableQuestionId"] not in previous_by_id
    ]
    chapter_counts = {
        source.chapter: sum(question["sourceChapter"] == source.chapter for question in questions)
        for source in SOURCES
    }
    return {
        "questionCount": len(questions),
        "chapterCounts": chapter_counts,
        "updatedQuestionCount": len(updated),
        "answerChangedQuestionCount": len(answer_changed),
        "addedQuestionCount": len(added),
        "removedQuestionCount": len(removed),
        "unmatchedQuestionCount": 0,
        "updatedQuestions": updated,
        "answerChangedQuestions": answer_changed,
        "addedQuestions": added,
        "removedQuestions": removed,
        "unmatchedQuestions": [],
    }


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def build(*, check_only: bool = False) -> dict[str, Any]:
    raw_questions = read_excel_sources()
    mapping = load_mapping()
    previous = load_previous_questions()
    questions, new_source_keys = assign_stable_ids(raw_questions, mapping)
    report = build_report(previous, questions, new_source_keys, mapping)
    version_input = [
        {key: value for key, value in question.items() if key != "contentRevision"}
        for question in sorted(questions, key=lambda item: item["stableQuestionId"])
    ]
    question_bank_version = digest(version_input)
    output = {
        "metadata": {
            "schemaVersion": 2,
            "databaseVersion": 1,
            "questionBankVersion": question_bank_version,
            "questionCount": len(questions),
            "source": "source-excel/*.xlsx",
            "chapters": [
                {
                    "id": source.chapter,
                    "title": source.title,
                    "filename": source.filename,
                    "questionCount": report["chapterCounts"][source.chapter],
                }
                for source in SOURCES
            ],
        },
        "questions": sorted(questions, key=lambda item: item["stableQuestionId"]),
    }
    report["questionBankVersion"] = question_bank_version

    if not check_only:
        write_json(OUTPUT_PATH, output)
        write_json(MAP_PATH, mapping)
        write_json(REPORT_PATH, report)
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="驗證 Excel 並產生無人機 App 題庫")
    parser.add_argument("--check", action="store_true", help="只驗證與計算，不寫入檔案")
    args = parser.parse_args()
    try:
        report = build(check_only=args.check)
    except ValidationErrors as exc:
        print("\n❌ Excel 題庫驗證失敗，未更新 questions.json，也不會部署：", file=sys.stderr)
        for error in exc.errors:
            print(f"- {error}", file=sys.stderr)
        return 1

    print("✅ Excel 題庫驗證成功")
    print(f"- 合計題數：{report['questionCount']}")
    print(f"- 修改題目：{report['updatedQuestionCount']}")
    print(f"- 答案變更：{report['answerChangedQuestionCount']}")
    print(f"- 新增題目：{report['addedQuestionCount']}")
    print(f"- 移除題目：{report['removedQuestionCount']}")
    print(f"- 題庫版本：{report['questionBankVersion']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
